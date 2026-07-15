import pytest
from docx.document import Document as doc
from docx import Document
from apps.word_doc_services.create_document_objects import (
    create_table,
    TableDefinition
)
 
def test_create_table_adds_table_to_document():
    # arrange
    test_doc = Document()

    # act
    create_table(TableDefinition.SINGLE_CELL_TEXT_BOX, test_doc)

    # assert
    assert test_doc.tables.count == 1


def test_create_table_handles_invalid_table_definition():
    # arrange
    tbl_definition: TableDefinition = TableDefinition(999)
    test_doc = Document()

    # act
    create_table(tbl_definition, test_doc)

    # assert
    assert test_doc.tables.count == 0


def test_create_table_adds_table_footer():
    # arrange
    test_doc_table_footer = "Test Table Footer"
    test_doc = Document()
    
    # act
    create_table(TableDefinition.SINGLE_CELL_TEXT_BOX, test_doc, test_doc_table_footer)
    test_para = test_doc.paragraphs[0]
    
    # assert
    assert test_doc.paragraphs.count == 1
    assert test_para.text == test_doc_table_footer
