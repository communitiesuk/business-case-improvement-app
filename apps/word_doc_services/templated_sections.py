from docx.document import Document
from docx.shared import Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE

from .formatting_methods import add_text, add_subheader, insert_bullet_point_list

from .table_definitions import TABLE_REGISTRY, TableDefinition

aptos_bold_font_name: str = "Aptos Bold"
aptos_font_name: str = "Aptos"

overview = (
    "1: Proposal participants and organisational details\n",
    "2: Proposal Summary"
)

sections = (
    "3: Strategic Case (4 questions)\n",
    "4: Economic Case (7 questions)\n",
    "5: Commercial Case (3 questions)\n",
    "6: Funding and Financial Case (4 questions)\n",
    "7: Management Case (5 questions)\n",
    "8: Other business cases"
)

approvals = (
    "9: Subject Matter Expert (SME) assurance, recommendations and declaration\n",
    "10: SRO/SCS sign off"
)

def create_what_youll_be_asked_section(doc: Document):
    p = doc.add_paragraph()

    what_youll_be_asked_run = p.add_run("What you'll be asked in this document:")
    what_youll_be_asked_run.font.size = Pt(16)
    what_youll_be_asked_run.font.bold = True
    what_youll_be_asked_run.font.name = "Aptos Bold"

    overview_paragraph = doc.add_paragraph()
    add_subheader(overview_paragraph, "Overview:\n")
    for item in overview:
        add_text(overview_paragraph, item)

    sections_paragraph = doc.add_paragraph()
    add_subheader(sections_paragraph, "Sections:\n")
    for item in sections:
        add_text(sections_paragraph, item)
    
    approvals_paragraph = doc.add_paragraph()
    add_subheader(approvals_paragraph, "Approvals:\n")
    for item in approvals:
        add_text(approvals_paragraph, item)

def create_before_you_start_section(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(45)
    header_text = p.add_run("Before you start:")
    header_text.font.size = Pt(16)
    header_text.font.bold = True
    header_text.font.name = "Aptos Bold"

    before_start_paragraph = doc.add_paragraph()
    add_text(before_start_paragraph, "This template follows the HM Treasury Green Book Five Case Model. The questions are designed to guide you through developing a clear, evidence-based business case.\n")
    
    insert_bullet_point_list(doc, [
        "Some questions may appear related, but each question collects different information. Take some time to read all questions in this template before you start drafting - it will help you avoid repetition.", 
        "Word counts are provided as a guide only. You do not need to use the full word count suggestions for each question."
        ]
    )

def create_table(table_type: TableDefinition, doc: Document):
    table_data = next((item for item in TABLE_REGISTRY if item.definition == table_type))

    if not table_data or not table_data.cells:
        print(f"Table data not found. Skipping. table_type: {table_type.name}")
        return

    # NOTE: if the last row(s) in a table should be blank, you need to add a blank record in the table def.
    # like this: _CellData(row=2, column=1, paragraphs=[ _ParagraphData()] the row will be taken and accounted for
    row_count = max(cell.row for cell in table_data.cells)
    column_count = max(cell.column for cell in table_data.cells)

    tbl = doc.add_table(row_count, column_count)
    tbl.style = "Table Grid" # adds gridlines to the table

    # resize the table column widths if they have been provided
    if table_data.column_widths:
        for col_idx, width in enumerate(table_data.column_widths):
            if col_idx < len(tbl.columns):
                for cell in tbl.columns[col_idx].cells:
                    cell.width = width

    # resize the table row heights if they have been provided
    if table_data.row_heights:
        for row_idx, height in enumerate(table_data.row_heights):
            if row_idx < len(tbl.rows):
                row = tbl.rows[row_idx]
                row.height = height
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    # populate each cell ith the corresponding data, including some default formatting
    for cell_data in table_data.cells:
        r_idx = cell_data.row - 1
        c_idx = cell_data.column - 1

        tbl_cell = tbl.cell(r_idx, c_idx)
        tbl_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        
        # add a paragrpah to the cell, then each run to that paragraph.
        # doing this allows for multiple formats within the same cell
        for p_idx, p_data in enumerate(cell_data.paragraphs):
            para = tbl_cell.paragraphs[0] if p_idx == 0 else tbl_cell.add_paragraph()

            for run_data in p_data.runs:
                run = para.add_run(run_data.text)
                run.italic = run_data.italic
                run.bold = run_data.bold
                run.font.name = aptos_font_name
                run.font.size = Pt(12)
