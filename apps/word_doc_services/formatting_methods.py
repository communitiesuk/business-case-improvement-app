
from docx.document import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt

aptos_bold_font_name: str = "Aptos Bold"
aptos_font_name: str = "Aptos"

def insert_bullet_point_list(doc: Document, items: list):
    for item in items:
        p = doc.add_paragraph()
        p.style="List Bullet"
        p.paragraph_format.left_indent = Pt(45)
        r = p.add_run(item)
        r.font.size = Pt(12)
        r.font.bold = False
        r.font.name = aptos_font_name

def add_subheader(p: Paragraph, content: str):
    r = p.add_run(content)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = aptos_bold_font_name

def add_text(p: Paragraph, content: str):
    r = p.add_run(content)
    r.font.size = Pt(12)
    r.font.bold = False
    r.font.name = aptos_font_name
