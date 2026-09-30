from docx.document import Document as doc
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from .table_common_resources import (
    table_width,
    _blue_help_box_background_hex
)

from word_doc_services.common_resources import (
    get_general_font_colour_rgb,
    _bold_font_name,
    _regular_font_name
)

'''
Summary:
    Exists to wrap the creation of the help box logic so as to not bloat the word doc wrapper.
    Creates the help box, formatting etc., and adds it to the Document
'''
class HelpBox:

    def __init__(self, contents: list[str]):
        self.contents = contents

    '''
    Summary:
        Add a help box the the document passed through.
    '''
    def add_help_box(self, doc: doc):
        p: Paragraph | None = None
        tbl = doc.add_table(1, 1)
        tbl._cells[0].width = table_width
        self.set_cell_background(tbl._cells[0])

        for idx, content in enumerate(self.contents):
            # the add_table() method above adds a paragraph by default, so use this
            # so the styling works with the indents and spaces as we use the default
            if idx == 0:
                p = tbl._cells[0].paragraphs[0]
            else:
                p = tbl._cells[0].add_paragraph()

            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p.paragraph_format.left_indent = Pt(8)
            p.paragraph_format.right_indent = Pt(8)
                            
            r = p.add_run()
            r.font.color.rgb = get_general_font_colour_rgb()
            r.font.size = Pt(12)

            if idx == 0:
                p.paragraph_format.space_before = Pt(8)
                r.font.bold = True
                r.font.name = _bold_font_name
                r.add_text(f"➡️ {content}")
            else:
                r.font.bold = False
                r.font.name = _regular_font_name
                r.add_text(content)

        # after the loop set the spacing for the last paragraph
        if p is not None:
            p.paragraph_format.space_after = Pt(8)


    '''Summary:
        Set the background to the blue colour we need for help boxes
    '''
    def set_cell_background(self, cell: _Cell):
        shading_elm_xml = f'<w:shd {nsdecls("w")} w:fill="{_blue_help_box_background_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_elm_xml))
