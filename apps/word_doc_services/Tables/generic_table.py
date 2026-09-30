from docx.document import Document as doc
from docx.shared import Pt, Cm

from apps.word_doc_services.tables.table_definitions import (
    TABLE_DEFINITION,
    TABLE_REGISTRY,
    _TableContent,
    HEADER_DIRECTION
)

from apps.word_doc_services.common_resources import _regular_font_name

from apps.word_doc_services.tables.table_common_resources import (
    max_table_width,
    set_default_paragraph_formatting
)

'''
Summary:
    Creates a generic table based on the table definition.
'''
class GenericTable():

    def __init__(self, tbl_def: TABLE_DEFINITION):
        self.tbl_def = tbl_def


    def create_table(self, doc: doc):
        tbl_content: _TableContent = next((item for item in TABLE_REGISTRY if item.definition == self.tbl_def))

        # determine direction of headers, and add that man rows or columns. Then add corresponding alternate value for alternate object
        rows_needed = len(tbl_content.headers) if tbl_content.header_direction == HEADER_DIRECTION.VERTICAL else tbl_content.headers_alternate_direction_object_count
        columns_needed = len(tbl_content.headers) if tbl_content.header_direction == HEADER_DIRECTION.HORIZONTAL else tbl_content.headers_alternate_direction_object_count

        tbl = doc.add_table(rows_needed, columns_needed, "Table Grid")
        
        # equal column widths
        col_width: float  = max_table_width / len(tbl.columns)

        header_row_index = 0
        header_column_index = 0

        for header in tbl_content.headers:
            try:
                p = tbl.cell(header_row_index, header_column_index).paragraphs[0]
                r = p.add_run(header)
                r.font.bold = tbl_content.bold_headers
                r.font.name = _regular_font_name
                r.font.size = Pt(12)
            except IndexError as ex:
                print(f"Error when creating the table for {self.tbl_def}. Err: {ex.__str__}")
                return

            if tbl_content.header_direction == HEADER_DIRECTION.VERTICAL:
                header_row_index += 1
            else:
                header_column_index += 1

        for content in tbl_content.extra_content:
            p = tbl.cell(content.row - 1, content.column - 1).paragraphs[0]
            r = p.add_run(content.content)
            r.italic = content.is_italic
            r.font.name = _regular_font_name
            r.font.size = Pt(12)

        # Widths have to be set per cell, not per column (a fun little gotcha).
        # Then for any empty cells we set the formatting so typing in it forces our desired format 
        for c in tbl._cells:
            c.width = Cm(col_width)
            if len(c.paragraphs[0].runs) == 0:
                set_default_paragraph_formatting(c.paragraphs[0])
