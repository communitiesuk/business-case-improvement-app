
from docx.document import Document as doc
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Cm, Pt

from .table_common_resources import (
    table_width,
    set_default_paragraph_formatting
)

from word_doc_services.common_resources import get_general_font_colour_rgb, _regular_font_name

class InputBox:

    def __init__(self, word_limit: str = ""):
        self.word_limit = word_limit


    '''
    Summary:
        Addds an input box to the 
    '''
    def add_input_box(self, doc: doc):
        tbl = doc.add_table(1, 1)
        box = tbl.rows[0]

        tbl._cells[0].width = table_width
        tbl.style = "Table Grid" # adds gridlines (in this case, a border) to the table
        box_paragraph = tbl._cells[0].paragraphs[0]
        set_default_paragraph_formatting(box_paragraph)

        box.height = Cm(2.55)
        box.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

        box_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        # add a footer to the input box with the word limit
        if self.word_limit != "":
            tbl_footer = doc.add_paragraph()
            tbl_footer.paragraph_format.space_before = 0

            r = tbl_footer.add_run("Word count guideline: {} words".format(self.word_limit))
            r.font.color.rgb = get_general_font_colour_rgb()
            r.font.size = Pt(12)
            r.font.bold = False
            r.font.italic = False
            r.font.name = _regular_font_name
