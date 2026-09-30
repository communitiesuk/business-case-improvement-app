from docx.document import Document as doc
from docx.shared import Cm, Pt

from table_common_resources import max_table_width
from word_doc_services.common_resources import _regular_font_name

'''
Summary:
    Class for creating a 'checkbox', a two-column table with the left hand column as a 'check' column,
    and the right hand side being an options column. This is neater and more user friendly than the checkbox control
    that is available in Word.
'''
class CheckBox():
    check_box_size: float = 1.3

    def __init__(self, options):
        self.options = options


    def add_checkbox(self, doc: doc):
        tbl = doc.add_table(len(self.options), 2, "Table Grid")
        
        for idx, opt in enumerate(self.options):
            p = tbl.cell(idx, 1).paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(opt)
            r.font.name = _regular_font_name
            r.font.size = Pt(12)

        for c in tbl.column_cells(0):
            c.width = Cm(self.check_box_size)

        for c in tbl.column_cells(1):
                    c.width = Cm(max_table_width.cm - self.check_box_size)

        for row in tbl.rows:
            row.height = Cm(self.check_box_size)
