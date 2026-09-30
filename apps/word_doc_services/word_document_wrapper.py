from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.text.run import Run

import docx.oxml.ns
import docx.opc.constants
import logging

from Tables.generic_table import GenericTable
from Tables.input_box import InputBox
from Tables.helpbox import HelpBox
from Tables.checkbox import CheckBox
from Tables.table_definitions import TABLE_DEFINITION

from common_resources import (
    get_general_font_colour_rgb,
    get_mhclg_green_rgb,
    _bold_font_name,
    _italic_font_name,
    _regular_font_name,
    _hyperlink_font_colour_hex
)

logger = logging.getLogger(__name__)

'''
Summary:
    Wrapper around the Word Document logic.
    
    Because Python-Docx uses proxy classes to handle the typing and what is returned, you can't have 
    objects like Paragraphs and Tables truly exist outside of the method to add_XXX() which returns the object.
    This requires having a doc somewhere which would couple things together.
    BusinessCaseWordDocumentWrapper is wrapping the logic so as to separate calling code fromm logic as much as possible.
    
    In this way the wrapper is souly responsible for styling/ handling the word document logic,
    and the methods are simply called by whatever needs a document created.
'''
class BusinessCaseWordDocumentWrapper:
    
    def __init__(self):
        self.doc = Document()


    def add_h1_section_header(self, header_text: str):
        p = self.get_basic_paragraph()
        r = p.add_run(header_text)
        r.font.color.rgb = get_mhclg_green_rgb()
        r.font.size = Pt(28)
        r.font.bold = True
        r.font.italic = False
        r.font.name = _bold_font_name


    def add_h2_section_header(self, header_text: str):
        p = self.get_basic_paragraph()
        r = p.add_run(header_text)
        r.font.color.rgb = get_mhclg_green_rgb()
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.italic = False
        r.font.name = _bold_font_name

    
    def add_paragraph(self, paragraph_content: str):
        p = self.get_basic_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        r = p.add_run(paragraph_content)
        self.style_paragraph_run(r)


    '''
    Summary:
        Get a paragraph with some basic styling applied that
        applies to all paragraphs.
    '''
    def get_basic_paragraph(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        return p


    '''
    Summary:
        Style a paragraph run.
        Separate styling here so multiple methods can call it.
        This is generic styling for a regualr paragraph, i.e not a header.
    '''
    def style_paragraph_run(self, r: Run):
        r.font.color.rgb = get_general_font_colour_rgb()
        r.font.size = Pt(12)
        r.font.bold = False
        r.font.italic = False
        r.font.name = _regular_font_name


    def add_bullet_point_list(self, items: list[str]):
        for item in items:
            p = self.doc.add_paragraph()
            p.style="List Bullet"
            p.paragraph_format.left_indent = Pt(45)
            r = p.add_run(item)
            r.font.color.rgb = get_general_font_colour_rgb()
            r.font.size = Pt(12)
            r.font.bold = False
            r.font.italic = False
            r.font.name = _regular_font_name


    '''
    Summary:
        Add fixed text (e.g. text from triage we know about) to the document.
    '''
    def add_fixed_text(self, fixed_text: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        r = p.add_run(fixed_text)
        r.font.color.rgb = get_general_font_colour_rgb()
        r.font.size = Pt(12)
        r.font.italic = True
        r.font.bold = False
        r.font.name = _italic_font_name


    '''
    Summary:
        Add a Hyperlink to the Word Document.
        Because there is no explicit hyperlink class in Docx we have to make it using the Oxml.
        Take in a string that represents the URL, the entire text to display, and the text that should become the hyperlink.
    '''
    def add_hyperlink(self, url: str, text: str, text_to_replace_with_hyperlink):
        if (text_to_replace_with_hyperlink not in text):
            logger.warning(f"Text to replace with a hyperlink does not exist in the paragraph string. Paragraph: {text}, Text to replace: {text_to_replace_with_hyperlink}")
            self.add_paragraph(text)
            return
        
        p = self.doc.add_paragraph()

        part = p.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

        # Create the w:hyperlink tag and add needed values
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(docx.oxml.ns.qn('r:id'), r_id, )

        # Create a w:r element
        new_run = OxmlElement('w:r')

        # Join all the xml elements together add add the required text to the w:r element
        new_run.append(self.get_hyperlink_run_style())
        new_run.text = text_to_replace_with_hyperlink
       
        sections = text.split(text_to_replace_with_hyperlink, 1)
        start = sections[0]
        end = sections[1]
        
        hyperlink.append(new_run)

        # join the text before the hyperlink, then the hyperlink, then the text after it
        start_run = p.add_run(f"{start.strip()} ")
        self.style_paragraph_run(start_run)

        p._p.append(hyperlink)

        end_run = p.add_run(f" {end.strip()}")
        self.style_paragraph_run(end_run)


    '''
    Summary:
        Style the hyperlink. Only the hyperlink itself, not the surrounding text.
    '''
    def get_hyperlink_run_style(self) -> BaseOxmlElement:
        # Create a new w:rPr element - Run Properties
        rPr = OxmlElement('w:rPr')
        
        color = OxmlElement('w:color')
        color.set(docx.oxml.ns.qn('w:val'), _hyperlink_font_colour_hex)
        rPr.append(color)

        underline = OxmlElement('w:u')
        underline.set(docx.oxml.ns.qn('w:val'), 'single')
        rPr.append(underline)

        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), _bold_font_name)
        rFonts.set(qn('w:hAnsi'), _bold_font_name)
        rPr.append(rFonts)

        # add Bold text
        b = OxmlElement('w:b')
        rPr.append(b)

        return rPr


    '''
    Summary:
        Add a help box to the Word Document.
    Params:
        contents: List of strings that the help box should contain.
                The first element will be in bold.
    '''
    def add_help_box(self, contents: list[str]):
        help_box = HelpBox(contents)
        help_box.add_help_box(self.doc)
        self.doc.add_paragraph()

    '''
    Summary:
        Add an input box to the Word Document.
    Params:
        word_limit: optional, guidance on how many words to use in the Input box.
    '''
    def add_input_box(self, word_limit: str):
        input_box = InputBox(word_limit)
        input_box.add_input_box(self.doc)
        self.doc.add_paragraph()

    '''
    Summary:
        Add a generic Table to the Word Document.
    Params:
        tbl_def: defines the table data to use.
    '''
    def add_generic_table(self, tbl_def: TABLE_DEFINITION):
        tbl = GenericTable(tbl_def)
        tbl.create_table(self.doc)
        self.doc.add_paragraph()

    '''
    Summary:
        Add a checkbox Table to the Word Document.
        To avoid using the in-built Word checkbox as that is clunky and awkward.
        This is preferred.
    '''
    def add_checkbox(self, options: list[str]):
        checkbox_tbl = CheckBox(options)
        checkbox_tbl.add_checkbox(self.doc)
        self.doc.add_paragraph()

    '''
    Summary;
        Save the document to the location required.
        If errors are encountered return a failing result.
    '''
    def save_document(self, save_location: str) -> bool:
        try:
            self.doc.save(save_location)
            return True
        except Exception as ex:
            logger.error(f"Error while saving Business Case template. Message: {ex.__str__}")
            return False
